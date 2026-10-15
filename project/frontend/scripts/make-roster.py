# docs/users.xlsx (1행 머리글: 팀 / 이름) -> src/data/roster.local.json
# 실행: python scripts/make-roster.py  (project/frontend 에서)
import json, sys
from pathlib import Path
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE.parents[2] / 'docs' / 'users.xlsx'
OUT = HERE.parent / 'src' / 'data' / 'roster.local.json'

clean = lambda v: ' '.join(str(v or '').split())
rows = list(load_workbook(SRC, read_only=True).active.iter_rows(values_only=True))
head = [clean(c) for c in rows[0]]
ti, ni = head.index('팀'), head.index('이름')
out, seen = [], set()
for r in rows[1:]:
    team, name = clean(r[ti]) or '미지정', clean(r[ni])
    if name and (team, name) not in seen:
        seen.add((team, name)); out.append({'team': team, 'name': name})
OUT.parent.mkdir(parents=True, exist_ok=True)
OUT.write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding='utf-8')
print(f'{len({m["team"] for m in out})}개 팀, {len(out)}명 -> {OUT}')
